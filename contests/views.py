import csv
import re
import io
import tempfile
import aspose.words as aw

from io import BytesIO
from xlsx2html import xlsx2html
from datetime import timedelta
from pygments import highlight
from pygments.formatters import HtmlFormatter
from pygments.lexers import CppLexer, TextLexer
from pydocx import PyDocX
from xls2xlsx import XLS2XLSX
from openpyxl import Workbook

try:
    from aspose import slides as slides
except ImportError:
    slides = None

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin, UserPassesTestMixin
from django.contrib.auth.models import User
from django.forms.models import inlineformset_factory, modelformset_factory
from django.http import Http404, HttpResponse, HttpResponseRedirect, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.views.generic import CreateView, DeleteView, DetailView, FormView, ListView, UpdateView
from django.views.generic.base import TemplateView, View
from django.views.generic.detail import BaseDetailView, SingleObjectMixin
from django.views.generic.edit import BaseUpdateView
from django.views.generic.list import BaseListView

from accounts.models import Account, Faculty, Notification, Announcement
from contest.mixins import LoginRedirectMixin, OwnershipOrMixin, LeadershipOrMixin, PaginatorMixin
from contest.soft_deletion import SoftDeletionUpdateView, SoftDeletionDeleteView
from contest.templatetags.views import has_leader_permission, get_updated_query_string
from contests.forms import (AssignmentForm, AssignmentEvaluateForm, AssignmentSetForm, AssignmentUpdateForm,
                            AssignmentUpdateAttachmentForm, ContestForm, ContestAttachmentForm, ContestMoveForm,
                            CourseFinishForm, CourseForm, CourseLeaderForm, CreditReportForm, CreditSetForm, FNTestForm,
                            OptionForm, OptionBaseFormSet, ProblemAttachmentForm, ProblemCommonForm, ProblemMoveForm,
                            ProblemProgramForm, ProblemRollbackResultsForm, ProblemTestForm, SubmissionFilesForm,
                            SubmissionMossForm, SubmissionOptionsForm, SubmissionPatternForm, SubmissionProgramForm,
                            SubmissionTextForm, SubmissionUpdateForm, SubmissionVerbalForm, UTTestForm, SubProblemForm)
from contests.models import (Assignment, Attachment, Contest, Course, CourseLeader, Credit, Execution, FNTest, Filter,
                             IOTest, Option, Problem, SubProblem, Submission, SubmissionPattern, UTTest)
from contests.results import TaskProgress
from contests.tasks import evaluate_submission, moss_submission
from contests.templatetags.contests import colorize

from schedule.models import Schedule

"""=================================================== Attachment ==================================================="""


class AttachmentDetail(DetailView):
    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        try:
            attachment = self.object.attachment_set.get(id=kwargs.get('attachment_id'))
        except Attachment.DoesNotExist:
            raise Http404("Attachment with id = %s does not exist." % kwargs.get('attachment_id'))
        if attachment.extension() not in ('.h', '.hpp',
                                          '.c', '.cpp',
                                          '.ppt', '.pptx',
                                          '.xls', '.xlsx',
                                          '.doc', '.docx',
                                          '.csv'):
            return HttpResponseRedirect(attachment.file.url)

        context = self.get_context_data(object=self.object, attachment=attachment)
        return self.render_to_response(context)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        attachment = kwargs.get('attachment')
        attachment_ext = attachment.extension()
        if attachment_ext in ('.h', '.hpp', '.c', '.cpp'):
            try:
                content = attachment.file.read()
            except FileNotFoundError:
                raise Http404("File %s does not exist." % attachment.filename)
            formatter = HtmlFormatter(linenos='inline', wrapcode=True)
            context['code'] = highlight(content.decode(errors='replace').replace('\t', ' ' * 4), CppLexer(), formatter)
        elif attachment_ext in ('.ppt', '.pptx') and slides is not None:
            pres = slides.Presentation(attachment.file.path)
            options = slides.export.HtmlOptions()
            out_stream = BytesIO()
            pres.save(out_stream, slides.export.SaveFormat.HTML, options)
            out_stream = str(out_stream.getvalue(), 'utf-8').replace('''Evaluation only.</tspan>''', " </tspan>") \
                .replace("Created with Aspose.Slides for .NET Standard 2.0 22.1.</tspan>", " </tspan>") \
                .replace("Copyright 2004-2022Aspose Pty Ltd.</tspan>", " </tspan>") \
                .replace('''class="slide"''', '''class="slide"  style="margin-bottom: 1%!important;"''') \
                .replace('''xlink="http://www.w3.org/1999/xlink" width=''',
                         '''xlink="http://www.w3.org/1999/xlink" class="mw-100 h-auto d-inline-block"''')
            context['code'] = out_stream.replace('''class="slideTitle"''', '''style="display:none;"''')
        elif attachment_ext in ('.xls', '.xlsx'):
            if attachment_ext == '.xls':
                temp = tempfile.TemporaryFile()
                converter = XLS2XLSX(attachment.file.path)
                converter.to_xlsx(temp)
            else:
                temp = attachment.file.path
            html_sheets = dict()
            r = re.compile(r'\<td id\=\".+\!.+\"')
            for i in range(100):
                try:
                    xlsx_file: str = temp
                    current_sheet = io.StringIO()
                    xlsx2html(xlsx_file, current_sheet, locale='en', sheet=i)
                    current_sheet.seek(0)
                    current_sheet_html = current_sheet.read()
                    sheet_name_str = r.search(current_sheet_html).group(0)
                    sheet_name = sheet_name_str[8:sheet_name_str.index('!')]
                    html_sheets[sheet_name] = current_sheet_html
                except IndexError:
                    break
            nav = '<br><ul>{}</ul>'.format(''.join(f'<li><a href="#fsheet{id(sheet_name)}">{sheet_name}</a></li>'
                                                   for sheet_name in html_sheets.keys()))
            context['code'] = ''.join(f'<hr><h3 id="fsheet{id(sheet_name)}">{sheet_name}</h3><hr>{value}'
                                      for sheet_name, value in html_sheets.items()) + nav
        elif attachment_ext == '.csv':
            temp = tempfile.TemporaryFile()
            workbook = Workbook()
            worksheet = workbook.active
            with open(attachment.file.path, 'r') as file:
                for row in csv.reader(file):
                    worksheet.append(row)
            workbook.save(temp)
            sheet = io.StringIO()
            xlsx2html(temp, sheet, locale='en')
            sheet.seek(0)
            sheet_html = sheet.read()
            context['code'] = sheet_html
        elif attachment_ext == '.doc':
            doc = aw.Document(attachment.file.path)
            out_stream = io.BytesIO()
            doc.save(out_stream, aw.SaveFormat.DOCX)
            # context['code'] = str(mammoth.convert_to_html(outStream_).value)\
            #     .replace('''Evaluation Only. Created with Aspose.Words. Copyright 2003-2022 Aspose Pty Ltd.''', " ")
            context['code'] = str(PyDocX.to_html(out_stream)).replace(
                '''Evaluation Only. Created with Aspose.Words. Copyright 2003-2022 Aspose Pty Ltd.''', " ")
        elif attachment_ext == '.docx':
            # with open(attachment.file.path, "rb") as docx_file:
            #    context['code'] = mammoth.convert_to_html(docx_file).value
            context['code'] = PyDocX.to_html(attachment.file.path)

        return context


class AttachmentDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = Attachment
    template_name = 'contests/attachment/attachment_delete.html'
    permission_required = 'contests.delete_attachment'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return hasattr(self.object.object, 'course') and self.object.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return hasattr(self.object.object, 'course') and self.object.object.course.leaders.filter(
            id=self.request.user.id).exists()

    def get_success_url(self):
        return self.object.object.get_absolute_url()


"""==================================================== Deleted ====================================================="""


class DeletedList(LoginRedirectMixin, PermissionRequiredMixin, TemplateView):
    template_name = 'contests/deleted/deleted_list.html'
    permission_required = 'contests.view_deleted'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['deleted_courses'] = Course.deleted_objects.all()
        context['deleted_contests'] = Contest.deleted_objects.all()
        context['deleted_problems'] = Problem.deleted_objects.all()
        return context


"""===================================================== Course ====================================================="""


class CourseDetail(LoginRedirectMixin, DetailView):
    model = Course
    template_name = 'contests/course/course_detail.html'


class CourseDiscussion(LoginRedirectMixin, PaginatorMixin, DetailView):
    model = Course
    template_name = 'contests/course/course_discussion.html'
    paginate_by = 30

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        comments = self.object.comment_set.actual()
        context['paginator'], context['page_obj'], context['comments'], context['is_paginated'] = \
            self.paginate_queryset(comments)
        return context


class CourseCreate(LoginRedirectMixin, PermissionRequiredMixin, CreateView):
    model = Course
    form_class = CourseForm
    template_name = 'contests/course/course_form.html'
    permission_required = 'contests.add_course'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated:
            faculty_id = int(request.GET.get('faculty_id') or request.user.account.faculty_id)
            self.storage['faculty'] = get_object_or_404(Faculty, id=faculty_id)
        return super().dispatch(request, *args, **kwargs)

    def get_initial(self):
        initial = super().get_initial()
        initial['faculty'] = self.storage['faculty']
        return initial

    def form_valid(self, form):
        form.instance.owner = self.request.user
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['title'] = "Добавление курса"
        return context


class CourseUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                   SoftDeletionUpdateView):
    model = Course
    form_class = CourseForm
    template_name = 'contests/course/course_form.html'
    permission_required = 'contests.change_course'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['title'] = "Редактирование курса"
        return context


class CourseDelete(LoginRedirectMixin, OwnershipOrMixin, PermissionRequiredMixin, SoftDeletionDeleteView):
    model = Course
    success_url = reverse_lazy('contests:course-list')
    template_name = 'contests/course/course_delete.html'
    permission_required = 'contests.delete_course'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.owner_id == self.request.user.id


class CourseList(LoginRedirectMixin, ListView):
    model = Course
    template_name = 'contests/course/course_list.html'
    context_object_name = 'courses'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated:
            faculty_id = int(request.GET.get('faculty_id') or request.user.account.faculty_id)
            self.storage['faculty'] = get_object_or_404(Faculty, id=faculty_id)
        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        return super().get_queryset().filter(faculty=self.storage['faculty'])

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.user.has_perm('accounts.add_faculty') or self.request.user.account.faculty.is_interfaculty:
            faculties = Faculty.objects.all()
        else:
            faculties = Faculty.objects.available_for_user(self.request.user)
        context['faculties'] = faculties
        context['faculty'] = self.storage['faculty']
        return context


"""================================================== CourseLeader =================================================="""


class CourseUpdateLeaders(LoginRedirectMixin, OwnershipOrMixin, PermissionRequiredMixin, FormView):
    template_name = 'contests/course_leader/course_leader_formset.html'
    permission_required = 'contests.change_course'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=kwargs.pop('pk'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def get_queryset(self):
        return CourseLeader.objects.filter(course=self.storage['course'])

    def get_form_class(self):
        extra = 0 if self.get_queryset().exists() else 1
        return modelformset_factory(CourseLeader, CourseLeaderForm, extra=extra, max_num=10, validate_max=True,
                                    can_delete=True)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['queryset'] = self.get_queryset()
        kwargs['form_kwargs'] = {
            'course': self.storage['course'],
            'leader_queryset': Account.staff.get_leader_queryset(self.storage['course'].faculty)
        }
        return kwargs

    def form_valid(self, form):
        self.objects = form.save()
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['formset'] = context['form']
        context.update(self.storage)
        return context

    def get_success_url(self):
        return reverse('contests:course-detail', kwargs={'pk': self.storage['course'].id})


"""===================================================== Credit ====================================================="""


class CreditReport(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, FormView):
    form_class = CreditReportForm
    template_name = 'contests/credit/credit_report.html'
    permission_required = 'contests.report_credit'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=kwargs.pop('course_id'))
        self.storage['faculty_id'] = int(request.GET.get('faculty_id') or 0)
        self.storage['debts'] = bool(request.GET.get('debts'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        course = self.storage['course']
        students = Account.students.enrolled()
        students = students.debtors(course) if self.storage['debts'] else students.current(course)
        examiners = Account.objects.filter(user__groups__name='Преподаватель')
        if course.faculty.short_name == "МФК":
            if self.storage['faculty_id'] > 0:
                students = students.filter(faculty_id=self.storage['faculty_id'])
                examiners = examiners.filter(faculty_id=self.storage['faculty_id'])
            students = students.order_by('faculty__short_name', 'user__last_name', 'user__first_name')
            examiners = examiners.order_by('faculty__short_name', 'user__last_name', 'user__first_name')
        else:
            students = students.order_by('user__last_name', 'user__first_name')
            examiners = examiners.filter(faculty=course.faculty).order_by('faculty__short_name', 'user__last_name',
                                                                          'user__first_name')
        kwargs['students'] = students
        kwargs['examiners'] = examiners
        kwargs['course'] = course
        return kwargs

    def post(self, request, *args, **kwargs):
        form = self.get_form()
        if form.is_valid():
            report_file, filename = Credit.objects.create_report(
                self.storage['course'], form.cleaned_data['group_name'], form.cleaned_data['students'],
                form.cleaned_data['type'], form.cleaned_data['examiners'], form.cleaned_data['faculty'],
                form.cleaned_data['discipline'], form.cleaned_data['semester'], form.cleaned_data['date']
            )
            response = HttpResponse(report_file, content_type='application/vnd.openxmlformats-officedocument')
            response['Content-Disposition'] = 'attachment; filename={}.docx'.format(filename)
            return response
        return self.form_invalid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.storage['course']
        return context


class CourseStart(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, FormView):
    form_class = CreditSetForm
    template_name = 'contests/course/course_start_form.html'
    permission_required = 'contests.add_credit'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=kwargs.pop('course_id'))
        self.storage['faculty_id'] = int(request.GET.get('faculty_id') or 0)
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        course = self.storage['course']
        students = Account.students.enrolled().allowed(course).filter(credit_id=None)
        if course.faculty.short_name == "МФК":
            if self.storage['faculty_id'] > 0:
                students = students.filter(faculty_id=self.storage['faculty_id'])
            students = students.order_by('faculty__short_name', '-level', 'user__last_name', 'user__first_name')
        else:
            students = students.order_by('-level', 'user__last_name', 'user__first_name')
        kwargs['runner_ups_queryset'] = students
        kwargs['course'] = course
        return kwargs

    def post(self, request, *args, **kwargs):
        form = self.get_form()
        if form.is_valid():
            Credit.objects.create_set(request.user, self.storage['course'], form.cleaned_data['runner_ups'])
            return self.form_valid(form)
        return self.form_invalid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.storage['course']
        return context

    def get_success_url(self):
        return reverse('contests:assignment-table', kwargs={'course_id': self.storage['course'].id})


class CourseFinish(LoginRedirectMixin, PermissionRequiredMixin, FormView):
    form_class = CourseFinishForm
    template_name = 'contests/course/course_finish_form.html'
    permission_required = 'accounts.change_account'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=kwargs.pop('course_id'))
        self.storage['faculty_id'] = int(request.GET.get('faculty_id') or 0)
        self.storage['debts'] = bool(request.GET.get('debts'))
        return super().dispatch(request, *args, **kwargs)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        course = self.storage['course']
        students = Account.students.enrolled()
        if self.storage['debts']:
            students = students.debtors(course)
        else:
            students = students.current(course)
        students = students.filter(credit_score__gte=1)
        if course.faculty.short_name == "МФК":
            if self.storage['faculty_id'] > 0:
                students = students.filter(faculty_id=self.storage['faculty_id'])
            students = students.order_by('faculty__short_name', '-level', 'user__last_name', 'user__first_name')
        else:
            students = students.order_by('-level', 'user__last_name', 'user__first_name')
        kwargs['level_ups_queryset'] = students
        kwargs['course'] = course
        return kwargs

    def post(self, request, *args, **kwargs):
        form = self.get_form()
        if form.is_valid():
            form.cleaned_data['level_ups'].level_up()
            return self.form_valid(form)
        return self.form_invalid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.storage['course']
        return context

    def get_success_url(self):
        return reverse('contests:assignment-table', kwargs={'course_id': self.storage['course'].id})


class CreditUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, UpdateView):
    model = Credit
    fields = ['score']
    template_name = 'contests/credit/credit_form.html'
    permission_required = 'contests.change_credit'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def form_valid(self, form):
        Notification.objects.notify(self.object.user, subject=self.request.user, action="изменил Вашу оценку по курсу",
                                    object=self.object.course)
        return redirect(self.get_success_url())

    def get_success_url(self):
        success_url = reverse('contests:assignment-table', kwargs={'course_id': self.object.course_id})
        return success_url + get_updated_query_string(self.request)


class CreditDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = Credit
    template_name = 'contests/credit/credit_delete.html'
    permission_required = 'contests.delete_credit'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:assignment-table', kwargs={'course_id': self.object.course_id})


"""===================================================== Filter ====================================================="""


class FilterCreateAPI(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, View):
    http_method_names = ['post']
    permission_required = 'contests.add_filter'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        course_id = request.POST.get('course_id')
        try:
            self.storage['course'] = Course.objects.get(id=course_id)
        except Course.DoesNotExist:
            return JsonResponse({'status': 'not_found'})
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def http_method_not_allowed(self, request, *args, **kwargs):
        return JsonResponse({'status': 'http_method_not_allowed'})

    def post(self, request, *args, **kwargs):
        response = {'status': 'ok'}
        _, created = Filter.objects.get_or_create(user_id=request.user.id, course=self.storage['course'])
        if not created:
            response['status'] = 'exists'
        return JsonResponse(response)

    def handle_no_permission(self):
        return JsonResponse({'status': 'access_denied'})


class FilterDeleteAPI(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, View):
    http_method_names = ['post']
    permission_required = 'contests.delete_filter'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        course_id = request.POST.get('course_id')
        try:
            self.storage['course'] = Course.objects.get(id=course_id)
        except Course.DoesNotExist:
            return JsonResponse({'status': 'not_found'})
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def http_method_not_allowed(self, request, *args, **kwargs):
        return JsonResponse({'status': 'http_method_not_allowed'})

    def post(self, request, *args, **kwargs):
        response = {'status': 'ok'}
        deleted, _ = Filter.objects.filter(user_id=request.user.id, course=self.storage['course']).delete()
        if not deleted > 0:
            response['status'] = 'not_found'
        return JsonResponse(response)

    def handle_no_permission(self):
        return JsonResponse({'status': 'access_denied'})


class FilterCreate(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, View):
    http_method_names = ['get']
    permission_required = 'contests.add_filter'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=request.GET.get('course_id'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get(self, request, *args, **kwargs):
        user_id = request.GET.get('user_id') or request.user.id
        Filter.objects.get_or_create(user_id=user_id, course=self.storage['course'])
        success_url = request.GET.get('from') or reverse('contests:filter-table')
        return HttpResponseRedirect(success_url)


class FilterDelete(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, View):
    http_method_names = ['get']
    permission_required = 'contests.delete_filter'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=request.GET.get('course_id'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get(self, request, *args, **kwargs):
        user_id = request.GET.get('user_id') or request.user.id
        Filter.objects.filter(user_id=user_id, course=self.storage['course']).delete()
        success_url = request.GET.get('from') or reverse('contests:filter-table')
        return HttpResponseRedirect(success_url)


class FilterTable(LoginRedirectMixin, PermissionRequiredMixin, TemplateView):
    template_name = 'contests/filter/filter_table.html'
    permission_required = 'contests.view_filter'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        users = User.objects.filter(groups__name__in=["Преподаватель", "Модератор"])
        courses = Course.objects.all()
        filters = list(Filter.objects.all().values_list('user_id', 'course_id'))
        table = []
        for i in range(len(users)):
            row = []
            for j in range(len(courses)):
                col = {'exists': (users[i].id, courses[j].id) in filters,
                       'params': {'user_id': users[i].id, 'course_id': courses[j].id}}
                row.append(col)
            table.append({'user': users[i], 'cols': row})
        context['users'] = users
        context['courses'] = courses
        context['table'] = table
        return context


"""===================================================== Contest ===================================================="""


class ContestDetail(LoginRedirectMixin, UserPassesTestMixin, DetailView):
    model = Contest
    template_name = 'contests/contest/contest_detail.html'

    def test_func(self):
        return not self.request.user.account.is_student or self.get_object().visible_to(self.request.user)


class ContestDiscussion(LoginRedirectMixin, UserPassesTestMixin, PaginatorMixin, DetailView):
    model = Contest
    template_name = 'contests/contest/contest_discussion.html'
    paginate_by = 30

    def test_func(self):
        return not self.request.user.account.is_student or self.get_object().visible_to(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        comments = self.object.comment_set.actual()
        context['paginator'], context['page_obj'], context['comments'], context['is_paginated'] = \
            self.paginate_queryset(comments)
        return context


class ContestAttachment(LoginRedirectMixin, UserPassesTestMixin, AttachmentDetail):
    model = Contest
    template_name = 'contests/contest/contest_attachment.html'

    def test_func(self):
        return not self.request.user.account.is_student or self.get_object().visible_to(self.request.user)


class ContestCreate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, CreateView):
    model = Contest
    form_class = ContestForm
    template_name = 'contests/contest/contest_form.html'
    permission_required = 'contests.add_contest'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=kwargs.pop('course_id'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get_initial(self):
        initial = super().get_initial()
        initial['course'] = self.storage['course']
        initial['number'] = Contest.objects.get_new_number(self.storage['course'])
        initial['title'] = "Раздел " + str(initial['number'])
        return initial

    def form_valid(self, form):
        form.instance.owner = self.request.user
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.storage['course']
        context['title'] = "Добавление раздела"
        return context


class ContestUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                    SoftDeletionUpdateView):
    model = Contest
    template_name = 'contests/contest/contest_form.html'
    permission_required = 'contests.change_contest'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['action'] = request.GET.get('action', None)
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_form_class(self):
        if self.storage['action'] == 'add_files':
            return ContestAttachmentForm
        elif self.storage['action'] == 'move':
            return ContestMoveForm
        else:
            return ContestForm

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.storage['action'] == 'move':
            kwargs['course_queryset'] = Course.objects.get_queryset_for_contest(self.object.course, self.request.user)
        return kwargs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.object.course
        if self.storage['action'] == 'add_files':
            context['title'] = "Добавление файлов в раздел"
        elif self.storage['action'] == 'move':
            context['title'] = "Перемещение раздела"
        elif self.storage['action'] == 'restore':
            context['title'] = "Восстановление раздела"
        else:
            context['title'] = "Редактирование раздела"
        return context


class ContestDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                    SoftDeletionDeleteView):
    model = Contest
    template_name = 'contests/contest/contest_delete.html'
    permission_required = 'contests.delete_contest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:course-detail', kwargs={'pk': self.object.course_id})


"""===================================================== Problem ===================================================="""


class ProblemDetail(LoginRedirectMixin, UserPassesTestMixin, PaginatorMixin, DetailView):
    model = Problem
    template_name = 'contests/problem/problem_detail.html'
    paginate_by = 10

    def test_func(self):
        return not self.request.user.account.is_student or self.get_object().visible_to(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        try:
            context['assignment'] = Assignment.objects.get(user=self.request.user, problem=self.object)
        except Assignment.DoesNotExist:
            context['assignment'] = None
        if self.request.user.has_perm('contests.view_submission_list') or has_leader_permission(self.request,
                                                                                                self.object.course):
            submissions = self.object.submission_set.all()
        else:
            submissions = self.object.submission_set.filter(owner_id=self.request.user.id)
        context['paginator'], context['page_obj'], context['submissions'], context['is_paginated'] = \
            self.paginate_queryset(submissions)
        if self.object.type == 'Test':
            context['subproblems'] = SubProblem.objects.filter(problem=self.object).select_related('sub_problem')
        return context


class ProblemDiscussion(LoginRedirectMixin, UserPassesTestMixin, PaginatorMixin, DetailView):
    model = Problem
    template_name = 'contests/problem/problem_discussion.html'
    paginate_by = 30

    def test_func(self):
        return not self.request.user.account.is_student or self.get_object().visible_to(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        comments = self.object.comment_set.actual()
        context['paginator'], context['page_obj'], context['comments'], context['is_paginated'] = \
            self.paginate_queryset(comments)
        return context


class ProblemAttachment(LoginRedirectMixin, UserPassesTestMixin, AttachmentDetail):
    model = Problem
    template_name = 'contests/problem/problem_attachment.html'

    def test_func(self):
        return not self.request.user.account.is_student or self.get_object().visible_to(self.request.user)


class ProblemRollbackResults(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                             FormView):
    form_class = ProblemRollbackResultsForm
    template_name = 'contests/problem/problem_rollback_results_form.html'
    permission_required = 'contests.change_problem'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['problem'] = get_object_or_404(Problem, id=self.kwargs.get('pk'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['problem'].course.owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['problem'].course.leaders.filter(id=self.request.user.id).exists()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['problem_id'] = self.kwargs.get('pk')
        return kwargs

    def post(self, request, *args, **kwargs):
        form = self.get_form()
        if form.is_valid():
            submissions = form.cleaned_data['submissions']
            assignments = Assignment.objects.to_rollback(submissions)
            submissions.rollback_status()
            assignments.rollback_score()
            return self.form_valid(form)
        else:
            return self.form_invalid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['problem'] = self.storage['problem']
        return context

    def get_success_url(self):
        return get_object_or_404(Problem, id=self.kwargs.get('pk')).get_absolute_url()


class ProblemCreate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, CreateView):
    model = Problem
    template_name = 'contests/problem/problem_form.html'
    permission_required = 'contests.add_problem'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['contest'] = get_object_or_404(Contest, id=kwargs.pop('contest_id'))
        self.storage['type'] = kwargs.get('type')
        if self.storage['type'] == 'Options':
            OptionFormSet = inlineformset_factory(parent_model=Problem, model=Option, form=OptionForm,
                                                  formset=OptionBaseFormSet, fields=('text', 'is_correct'), extra=0,
                                                  min_num=2, max_num=20, validate_min=True, validate_max=True)
            if request.method == 'GET':
                formset = OptionFormSet()
            else:
                formset = OptionFormSet(data=request.POST, files=request.FILES)
            self.storage['formset'] = formset
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['contest'].course.owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['contest'].course.leaders.filter(id=self.request.user.id).exists()

    def get_form_class(self):
        if self.storage['type'] == 'Program':
            return ProblemProgramForm
        elif self.storage['type'] == 'Test':
            return ProblemTestForm
        else:
            return ProblemCommonForm

    def get_initial(self):
        initial = super().get_initial()
        initial['contest'] = self.storage['contest']
        initial['type'] = self.storage['type']
        initial['number'] = Problem.objects.get_new_number(self.storage['contest'])
        if self.storage['type'] == 'Program':
            initial['title'] = "Задача "
        elif self.storage['type'] == 'Test':
            initial['title'] = "Тест "
        else:
            initial['title'] = "Вопрос "
        initial['title'] += str(initial['number'])
        return initial

    def form_valid(self, form):
        form.instance.owner = self.request.user
        if self.storage['type'] == 'Options':
            formset = self.storage.get('formset')
            if formset.is_valid():
                self.object = form.save()
                formset.instance = self.object
                formset.save()
                return HttpResponseRedirect(self.get_success_url())
            else:
                self.storage['formset'] = formset
                return self.form_invalid(form)
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['contest'] = self.storage['contest']
        context['type'] = self.storage['type']
        context['formset'] = self.storage.get('formset')
        context['title'] = "Добавление задачи"
        return context

    def get_success_url(self):
        return reverse('contests:contest-detail', kwargs={'pk': self.object.contest_id})


class ProblemUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                    SoftDeletionUpdateView):
    model = Problem
    template_name = 'contests/problem/problem_form.html'
    permission_required = 'contests.change_problem'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['action'] = request.GET.get('action', None)
        problem = self.get_object()
        if problem.type == 'Options':
            OptionFormSet = inlineformset_factory(parent_model=Problem, model=Option, form=OptionForm,
                                                  formset=OptionBaseFormSet, fields=('text', 'is_correct'), extra=0,
                                                  min_num=2, max_num=20, validate_min=True, validate_max=True)
            if request.method == 'GET':
                formset = OptionFormSet(instance=problem)
            else:
                formset = OptionFormSet(data=request.POST, files=request.FILES, instance=problem)
            self.storage['formset'] = formset
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_form_class(self):
        if self.storage['action'] == 'add_files':
            return ProblemAttachmentForm
        elif self.storage['action'] == 'move':
            return ProblemMoveForm
        elif self.object.type == 'Program':
            return ProblemProgramForm
        elif self.object.type == 'Test':
            return ProblemTestForm
        else:
            return ProblemCommonForm

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.storage['action'] == 'move':
            kwargs['contest_queryset'] = Contest.objects.get_queryset_for_problem(self.object.contest,
                                                                                  self.request.user)
        return kwargs

    def form_valid(self, form):
        if self.object.type == 'Options':
            formset = self.storage.get('formset')
            if formset.is_valid():
                self.object = form.save()
                formset.save()
                return HttpResponseRedirect(self.get_success_url())
            else:
                self.storage['formset'] = formset
                return self.form_invalid(form)
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['contest'] = self.object.contest
        context['type'] = self.object.type
        context['formset'] = self.storage.get('formset')
        if self.storage['action'] == 'add_files':
            context['title'] = "Добавление файлов в задачу"
        elif self.storage['action'] == 'move':
            context['title'] = "Перемещение задачи"
        elif self.storage['action'] == 'restore':
            context['title'] = "Восстановление задачи"
        else:
            context['title'] = "Редактирование задачи"
        return context


class ProblemDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                    SoftDeletionDeleteView):
    model = Problem
    template_name = 'contests/problem/problem_delete.html'
    permission_required = 'contests.delete_problem'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:contest-detail', kwargs={'pk': self.object.contest_id})


"""=================================================== SubProblem ==================================================="""


class SubProblemUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, UpdateView):
    model = SubProblem
    form_class = SubProblemForm
    template_name = 'contests/subproblem/subproblem_form.html'
    permission_required = 'contests.change_subproblem'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:problem-detail', kwargs={'pk': self.object.problem_id})


class SubProblemDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = SubProblem
    template_name = 'contests/subproblem/subproblem_delete.html'
    permission_required = 'contests.delete_subproblem'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:problem-detail', kwargs={'pk': self.object.problem_id})


"""=============================================== SubmissionPattern ================================================"""


class SubmissionPatternDetail(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                              DetailView):
    model = SubmissionPattern
    context_object_name = 'submission_pattern'
    template_name = 'contests/submission_pattern/submission_pattern_detail.html'
    permission_required = 'contests.view_submission_pattern'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.object.problems.first().course
        return context


class SubmissionPatternCreate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                              CreateView):
    model = SubmissionPattern
    form_class = SubmissionPatternForm
    template_name = 'contests/submission_pattern/submission_pattern_form.html'
    permission_required = 'contests.add_submission_pattern'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['problem'] = get_object_or_404(Problem, id=kwargs.pop('problem_id'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['problem'].course.owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['problem'].course.leaders.filter(id=self.request.user.id).exists()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['problem'] = self.storage['problem']
        return kwargs

    def form_valid(self, form):
        form.instance.owner = self.request.user
        form.instance.problem = self.storage['problem']
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['problem'] = self.storage['problem']
        return context


class SubmissionPatternUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                              UpdateView):
    model = SubmissionPattern
    form_class = SubmissionPatternForm
    context_object_name = 'submission_pattern'
    template_name = 'contests/submission_pattern/submission_pattern_form.html'
    permission_required = 'contests.change_submission_pattern'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.leaders.filter(id=self.request.user.id).exists()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['problem'] = self.object.problems.first()
        return kwargs


class SubmissionPatternDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                              DeleteView):
    model = SubmissionPattern
    context_object_name = 'submission_pattern'
    template_name = 'contests/submission_pattern/submission_pattern_delete.html'
    permission_required = 'contests.delete_submission_pattern'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:submission-pattern-detail', kwargs={'pk': self.object.problem_id})


"""===================================================== IOTest ====================================================="""


class IOTestDetail(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DetailView):
    model = IOTest
    template_name = 'contests/iotest/iotest_detail.html'
    permission_required = 'contests.view_iotest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()


class IOTestCreate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, CreateView):
    model = IOTest
    fields = ['title', 'compile_args', 'compile_args_override', 'launch_args', 'launch_args_override', 'input',
              'output']
    template_name = 'contests/iotest/iotest_form.html'
    permission_required = 'contests.add_iotest'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['problem'] = get_object_or_404(Problem, id=kwargs.pop('problem_id'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['problem'].course.owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['problem'].course.leaders.filter(id=self.request.user.id).exists()

    def form_valid(self, form):
        form.instance.owner = self.request.user
        form.instance.problem = self.storage['problem']
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['problem'] = self.storage['problem']
        return context

    def get_success_url(self):
        if Submission.objects.to_rollback(self.object.problem_id).exists():
            return reverse('contests:problem-rollback-results', kwargs={'pk': self.object.problem_id})
        else:
            return super().get_success_url()


class IOTestUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, UpdateView):
    model = IOTest
    fields = ['title', 'compile_args', 'compile_args_override', 'launch_args', 'launch_args_override', 'input',
              'output']
    template_name = 'contests/iotest/iotest_form.html'
    permission_required = 'contests.change_iotest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['problem'] = self.object.problem
        return context

    def get_success_url(self):
        if Submission.objects.to_rollback(self.object.problem_id).exists():
            return reverse('contests:problem-rollback-results', kwargs={'pk': self.object.problem_id})
        else:
            return super().get_success_url()


class IOTestDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = IOTest
    template_name = 'contests/iotest/iotest_delete.html'
    permission_required = 'contests.delete_iotest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:problem-detail', kwargs={'pk': self.object.problem_id})


"""===================================================== UTTest ====================================================="""


class UTTestDetail(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DetailView):
    model = UTTest
    template_name = 'contests/uttest/uttest_detail.html'
    permission_required = 'contests.view_uttest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()


class UTTestAttachment(LoginRedirectMixin, AttachmentDetail):
    model = UTTest
    template_name = 'contests/uttest/uttest_attachment.html'


class UTTestCreate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, CreateView):
    model = UTTest
    form_class = UTTestForm
    template_name = 'contests/uttest/uttest_form.html'
    permission_required = 'contests.add_uttest'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['problem'] = get_object_or_404(Problem, id=kwargs.pop('problem_id'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['problem'].course.owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['problem'].course.leaders.filter(id=self.request.user.id).exists()

    def form_valid(self, form):
        form.instance.owner = self.request.user
        form.instance.problem = self.storage['problem']
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['problem'] = self.storage['problem']
        return context

    def get_success_url(self):
        if Submission.objects.to_rollback(self.object.problem_id).exists():
            return reverse('contests:problem-rollback-results', kwargs={'pk': self.object.problem_id})
        else:
            return super().get_success_url()


class UTTestUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, UpdateView):
    model = UTTest
    form_class = UTTestForm
    template_name = 'contests/uttest/uttest_form.html'
    permission_required = 'contests.change_uttest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['problem'] = self.object.problem
        return context

    def get_success_url(self):
        if Submission.objects.to_rollback(self.object.problem_id).exists():
            return reverse('contests:problem-rollback-results', kwargs={'pk': self.object.problem_id})
        else:
            return super().get_success_url()


class UTTestDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = UTTest
    template_name = 'contests/uttest/uttest_delete.html'
    permission_required = 'contests.delete_uttest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problem.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:problem-detail', kwargs={'pk': self.object.problem_id})


"""===================================================== FNTest ====================================================="""


class FNTestDetail(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DetailView):
    model = FNTest
    template_name = 'contests/fntest/fntest_detail.html'
    permission_required = 'contests.view_fntest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.object.problems.first().course
        return context


class FNTestCreate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, CreateView):
    model = FNTest
    form_class = FNTestForm
    template_name = 'contests/fntest/fntest_form.html'
    permission_required = 'contests.add_fntest'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['problem'] = get_object_or_404(Problem, id=kwargs.pop('problem_id'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['problem'].course.owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['problem'].course.leaders.filter(id=self.request.user.id).exists()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['problem'] = self.storage['problem']
        return kwargs

    def form_valid(self, form):
        form.instance.owner = self.request.user
        form.instance.problem = self.storage['problem']
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['problem'] = self.storage['problem']
        return context


class FNTestUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, UpdateView):
    model = FNTest
    form_class = FNTestForm
    template_name = 'contests/fntest/fntest_form.html'
    permission_required = 'contests.change_fntest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.leaders.filter(id=self.request.user.id).exists()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['problem'] = self.object.problems.first()
        return kwargs


class FNTestDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = FNTest
    template_name = 'contests/fntest/fntest_delete.html'
    permission_required = 'contests.delete_fntest'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.problems.first().course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:problem-detail', kwargs={'pk': self.object.problem_id})


"""=================================================== Assignment ==================================================="""


class AssignmentDetail(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, PaginatorMixin,
                       DetailView):
    model = Assignment
    template_name = 'contests/assignment/assignment_detail.html'
    permission_required = 'contests.view_assignment'
    paginate_by = 10

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id or self.object.user_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        submissions = self.object.get_submissions()
        context['submission_paginator'], context['submission_page_obj'], context['submissions'], context[
            'submission_is_paginated'] = \
            self.paginate_queryset(submissions)
        comments = self.object.comment_set.actual()
        context['comment_paginator'], context['comment_page_obj'], context['comments'], context[
            'comment_is_paginated'] = \
            self.paginate_queryset(comments)
        context['current_time'] = timezone.now()
        return context


class AssignmentDiscussion(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                           PaginatorMixin, DetailView):
    model = Assignment
    template_name = 'contests/assignment/assignment_discussion.html'
    permission_required = 'contests.view_assignment'
    paginate_by = 30

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id or self.object.user_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        comments = self.object.comment_set.actual()
        context['paginator'], context['page_obj'], context['comments'], context['is_paginated'] = \
            self.paginate_queryset(comments)
        return context


class AssignmentAttachment(LoginRedirectMixin, AttachmentDetail):
    model = Assignment
    template_name = 'contests/assignment/assignment_attachment.html'


class AssignmentCreate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, CreateView):
    model = Assignment
    form_class = AssignmentForm
    template_name = 'contests/assignment/assignment_form.html'
    permission_required = 'contests.add_assignment'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['course'] = get_object_or_404(Course, id=kwargs.pop('course_id'))
        self.storage['debts'] = bool(request.GET.get('debts'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get(self, request, *args, **kwargs):
        contest_id, user_id = request.GET.get('contest_id'), request.GET.get('user_id')
        if contest_id and user_id:
            self.storage['user'] = get_object_or_404(User, id=user_id)
            self.storage['contest'] = get_object_or_404(Contest, id=contest_id)
        return super().get(request, *args, **kwargs)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs.update(self.storage)
        return kwargs

    def form_valid(self, form):
        form.instance.owner = self.request.user
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.storage['course']
        context['title'] = "Добавление задания"
        return context

    def get_success_url(self):
        success_url = reverse('contests:assignment-table', kwargs={'course_id': self.storage['course'].id})
        return success_url + get_updated_query_string(self.request)


class AssignmentCreateRandomSet(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                                FormView):
    form_class = AssignmentSetForm
    template_name = 'contests/assignment/assignment_random_set_form.html'
    permission_required = 'contests.add_assignment'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated:
            course = get_object_or_404(Course, id=kwargs.pop('course_id'))
            user_as_leader_queryset = CourseLeader.objects.filter(course=course, leader=request.user)
            user_as_leader = user_as_leader_queryset.get() if user_as_leader_queryset.exists() else None
            default_faculty_id = 0
            if course.faculty.is_interfaculty and user_as_leader is not None:
                default_faculty_id = user_as_leader.leader.account.faculty_id
            default_group = user_as_leader.group if user_as_leader is not None else 0
            default_subgroup = user_as_leader.subgroup if user_as_leader is not None else 0
            self.storage['course'] = course
            self.storage['faculty_id'] = int(request.GET.get('faculty_id') or default_faculty_id)
            self.storage['group'] = int(request.GET.get('group') or default_group)
            self.storage['subgroup'] = int(request.GET.get('subgroup') or default_subgroup)
            self.storage['debts'] = bool(request.GET.get('debts'))
            contest_id = int(request.GET.get('contest_id') or 0)
            if contest_id:
                self.storage['contest'] = get_object_or_404(Contest, id=contest_id, course=self.storage['course'])
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get_initial(self):
        initial = super().get_initial()
        if 'contest' in self.storage:
            initial['contest'] = self.storage['contest']
        return initial

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['course'] = self.storage['course']
        return kwargs

    def post(self, request, *args, **kwargs):
        form = self.get_form()
        if form.is_valid():
            Assignment.objects.create_random_set(request.user, form.cleaned_data['contest'], form.cleaned_data['type'],
                                                 form.cleaned_data['limit_per_user'],
                                                 form.cleaned_data['submission_limit'], form.cleaned_data['deadline'],
                                                 self.storage)
            return self.form_valid(form)
        return self.form_invalid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.storage['course']
        return context

    def get_success_url(self):
        success_url = reverse('contests:assignment-table', kwargs={'course_id': self.storage['course'].id})
        return success_url + get_updated_query_string(self.request)


class AssignmentUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, UpdateView):
    model = Assignment
    template_name = 'contests/assignment/assignment_form.html'
    permission_required = 'contests.change_assignment'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['action'] = request.GET.get('action', None)
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def form_valid(self, form):
        Notification.objects.notify(self.object.user, subject=self.request.user, action="изменил Вашу оценку к задаче",
                                    object=self.object, relation="из раздела", reference=self.object.contest)
        return redirect('contests:assignment-detail', self.object.id)

    def get_form_class(self):
        if self.storage['action'] == 'add_files':
            return AssignmentUpdateAttachmentForm
        elif self.storage['action'] == 'evaluate':
            return AssignmentEvaluateForm
        else:
            return AssignmentUpdateForm

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.storage['action'] is None:
            kwargs['course'] = self.object.course
        return kwargs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.object.course
        context['action'] = self.storage['action']
        if self.storage['action'] == 'add_files':
            context['title'] = "Прикрепление файлов к заданию"
        elif self.storage['action'] == 'evaluate':
            context['title'] = "Оценка задания"
        else:
            context['title'] = "Редактирование задания"
        return context


class AssignmentDelete(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = Assignment
    template_name = 'contests/assignment/assignment_delete.html'
    permission_required = 'contests.delete_assignment'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_success_url(self):
        return reverse('contests:assignment-table', kwargs={'course_id': self.object.course.id})


class AssignmentUserTable(LoginRedirectMixin, ListView):
    model = Assignment
    template_name = 'contests/assignment/assignment_user_base.html'
    context_object_name = 'assignments'

    def get_queryset(self):
        return (Assignment.objects
                .filter(user=self.request.user)
                .select_related('problem', 'problem__contest')
                .order_by('-problem__contest__course', '-problem__contest', '-date_created'))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['credits'] = self.request.user.credit_set.select_related('course')
        context['notifications'] = Notification.objects.get_queryset().filter(recipient=self.request.user).unread()[:10]
        context['schedules'] = Schedule.objects.get_queryset().filter(
            date_from__range=[timezone.now().date() - timedelta(days=6), timezone.now().date() + timedelta(days=7)])
        context['announcements'] = Announcement.objects.get_queryset().filter(
            date_created__range=[timezone.now().date() - timedelta(days=90), timezone.now().date()])
        context['count_of_news'] = context['notifications'].count() + context['schedules'].count() + context[
            'announcements'].count()
        context['has_unread_notifications'] = Notification.objects.get_queryset().filter(
            recipient=self.request.user).unread().exists()
        return context


class AssignmentCourseTable(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, ListView):
    model = Assignment
    template_name = 'contests/assignment/assignment_course_base.html'
    context_object_name = 'assignments'
    permission_required = 'contests.view_assignment_table'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated:
            course = get_object_or_404(Course, id=kwargs.pop('course_id'))
            user_as_leader_queryset = CourseLeader.objects.filter(course=course, leader=request.user)
            user_as_leader = user_as_leader_queryset.get() if user_as_leader_queryset.exists() else None
            default_faculty_id = 0
            if course.faculty.is_interfaculty and user_as_leader is not None:
                default_faculty_id = user_as_leader.leader.account.faculty_id
            default_group = user_as_leader.group if user_as_leader is not None else 0
            default_subgroup = user_as_leader.subgroup if user_as_leader is not None else 0
            self.storage['course'] = course
            self.storage['faculty_id'] = int(request.GET.get('faculty_id') or default_faculty_id)
            self.storage['group'] = int(request.GET.get('group') or default_group)
            self.storage['subgroup'] = int(request.GET.get('subgroup') or default_subgroup)
            self.storage['debts'] = int(request.GET.get('debts') or 0)
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get_queryset(self):
        course = self.storage['course']
        if self.storage['debts']:
            students = Account.students.enrolled().debtors(course)
        else:
            students = Account.students.enrolled().current(course)
            if self.storage['faculty_id'] > 0:
                students = students.filter(faculty_id=self.storage['faculty_id'])
            if self.storage['group'] > 0:
                students = students.filter(group=self.storage['group'])
            if self.storage['subgroup'] > 0:
                students = students.filter(subgroup=self.storage['subgroup'])
        self.storage['students'] = students
        bool(self.storage['students'])  # evaluate now
        return Assignment.objects.for_course_table(course, self.storage['students'])

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['faculties'] = Faculty.objects.exclude(short_name="МФК")
        context.update(self.storage)
        return context


"""=================================================== Submission ==================================================="""


class SubmissionDetail(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, PaginatorMixin,
                       UpdateView):
    model = Submission
    form_class = SubmissionUpdateForm
    sub_form_class = SubmissionUpdateForm
    template_name = 'contests/submission/submission_detail.html'
    permission_required = 'contests.view_submission'
    paginate_by = 30

    def __init__(self):
        super().__init__()
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.object = self.get_object()
        if self.object.problem.type == 'Test':
            sub_submission_id = request.POST.get('sub_submission_id')
            if sub_submission_id is not None:
                if sub_submission_id.isdigit():
                    sub_submission_id = int(sub_submission_id)
                    self.storage['sub_submission_id'] = sub_submission_id
                    if not self.object.sub_submissions.filter(id=sub_submission_id).exists():
                        raise Http404("Sub submission does not exist")
                else:
                    raise Http404("sub_submission_id must contain only digits")
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id or self.object.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_form(self, form_class=None):
        sub_form_class = self.sub_form_class
        if self.object.problem.type == 'Test':
            forms = dict()
            sub_submission_id = self.storage.get('sub_submission_id')
            for sub_submission in self.object.sub_submissions.order_by('id'):
                if sub_submission.id == sub_submission_id:
                    form = sub_form_class(instance=sub_submission, data=self.request.POST)
                    forms[sub_submission.id] = self.storage['sub_form'] = form
                else:
                    forms[sub_submission.id] = sub_form_class(instance=sub_submission)
            self.storage['forms'] = forms
        form_class = self.get_form_class()
        if 'sub_form' not in self.storage:
            return super().get_form()  # filled with POST data
        else:
            return form_class(instance=self.object)  # not filled with data

    def post(self, request, *args, **kwargs):
        form = self.get_form()
        sub_form = self.storage.get('sub_form')
        if sub_form:
            if sub_form.is_valid():
                sub_form.save()
                return HttpResponseRedirect(self.get_success_url())
            return super().form_invalid(form)
        if form.is_valid():
            return super().form_valid(form)
        else:
            return super().form_invalid(form)

    def get(self, request, *args, **kwargs):
        if not hasattr(self, 'object'):  # self.object may be set in LoginRedirectOwnershipOrPermissionRequiredMixin
            self.object = self.get_object()
        # TODO: mark corresponding notification as read
        # Notification.objects.filter(recipient=request.user,
        #                             object_type=ContentType.objects.get_for_model(self.object),
        #                             object_id=self.object.id).mark_as_read()
        context = self.get_context_data(object=self.object)
        return self.render_to_response(context)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        comments = self.object.comment_set.actual()
        context['paginator'], context['page_obj'], context['comments'], context['is_paginated'] = \
            self.paginate_queryset(comments)
        context['current_time'] = timezone.now()
        context['from_assignment'] = 'from_assignment' in self.request.GET
        if 'forms' in self.storage:
            context['forms'] = self.storage['forms'].values()
        return context

    def get_success_url(self):
        success_url = self.object.get_absolute_url()
        if 'from_assignment' in self.request.GET:
            success_url += '?from_assignment=1'
        return success_url


class SubmissionDownload(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                         BaseDetailView):
    model = Submission
    permission_required = 'contests.download_submission'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        try:
            zip_file = self.object.get_files_as_zip()
        except FileNotFoundError:
            raise Http404("Cannot create zip file: some files are missing")
        response = HttpResponse(zip_file, content_type='application/zip')
        response['Content-Disposition'] = 'attachment; filename={}.zip'.format(self.object.pk)
        return response


class SubmissionAttachment(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                           AttachmentDetail):
    model = Submission
    template_name = 'contests/submission/submission_attachment.html'
    permission_required = 'contests.view_submission'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['from_assignment'] = 'from_assignment' in self.request.GET
        return context


class SubmissionCreate(LoginRedirectMixin, PermissionRequiredMixin, CreateView):
    model = Submission
    template_name = 'contests/submission/submission_form.html'
    permission_required = 'contests.add_submission'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated:
            problem = get_object_or_404(Problem, id=kwargs.pop('problem_id'))
            try:
                self.storage['assignment'] = Assignment.objects.get(user=self.request.user, problem=problem)
            except Assignment.DoesNotExist:
                self.storage['assignment'] = None
            if problem.type == 'Test':
                self.storage['main_problem'] = problem
                sub_problem_id = request.GET.get("sub_problem")
                main_submission_id = kwargs.pop('submission_id', None)
                sub_problem = None
                if main_submission_id:
                    main_submission = get_object_or_404(Submission, id=main_submission_id)
                    self.storage['main_submission'] = main_submission
                    pending_sub_problems = problem.sub_problems.exclude(
                        submission__in=main_submission.sub_submissions.all())
                    if sub_problem_id and sub_problem_id.isdigit():
                        sub_problem = pending_sub_problems.filter(id=sub_problem_id).first()
                    if sub_problem is None:
                        sub_problem = pending_sub_problems.first()
                    if sub_problem is None:
                        self.storage['main_submission_complete'] = True
                        return HttpResponseRedirect(self.get_success_url())
                    self.storage['problem'] = sub_problem
                else:
                    if sub_problem_id and sub_problem_id.isdigit():
                        sub_problem = problem.sub_problems.filter(id=sub_problem_id).first()
                    if sub_problem is None:
                        sub_problem = problem.sub_problems.first()
                    self.storage['problem'] = sub_problem
            else:
                self.storage['problem'] = problem
        return super().dispatch(request, *args, **kwargs)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['owner'] = self.request.user
        kwargs['problem'] = self.storage['problem']
        kwargs['assignment'] = self.storage['assignment']
        if 'main_submission' in self.storage:
            kwargs['main_submission'] = self.storage['main_submission']
        return kwargs

    def get_form_class(self):
        problem = self.storage['problem']
        if problem.type == 'Text':
            return SubmissionTextForm
        elif problem.type == 'Verbal':
            return SubmissionVerbalForm
        elif problem.type == 'Files':
            return SubmissionFilesForm
        elif problem.type == 'Options':
            return SubmissionOptionsForm
        else:
            return SubmissionProgramForm

    def form_valid(self, form):
        form.instance.owner = self.request.user
        form.instance.problem = self.storage['problem']
        if 'main_problem' in self.storage:
            if 'main_submission' not in self.storage:
                self.storage['main_submission'] = Submission.objects.create(problem=self.storage.get('main_problem'),
                                                                            owner=self.request.user,
                                                                            assignment=self.storage['assignment'])
        else:
            form.instance.assignment = self.storage['assignment']
        form.instance.main_submission = self.storage.get('main_submission')
        self.object = form.save()
        if self.object.problem.type == 'Program' and self.object.problem.is_testable:
            task = evaluate_submission.delay(self.object.pk, self.request.user.id)
            self.object.task_id = task.id
            self.object.save()
        if self.object.problem.type == 'Options':
            self.object.update_options_score()
        return HttpResponseRedirect(self.get_success_url())

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        problem = self.storage['problem']
        context['problem'] = problem
        context['assignment'] = self.storage['assignment']
        main_problem = self.storage.get('main_problem')
        if main_problem:
            context['main_problem'] = main_problem
            main_submission = self.storage.get('main_submission')
            if main_submission:
                context['main_submission'] = main_submission
                context['pending_sub_problems'] = main_problem.sub_problems.exclude(
                    submission__in=main_submission.sub_submissions.all()).exclude(id=problem.id)
            else:
                context['pending_sub_problems'] = main_problem.sub_problems.exclude(id=problem.id)
        context['current_time'] = timezone.now()
        context['from_assignment'] = 'from_assignment' in self.request.GET
        return context

    def get_success_url(self):
        if 'main_submission' in self.storage:
            main_submission = self.storage['main_submission']
            if 'main_submission_complete' in self.storage:
                success_url = reverse('contests:submission-detail', kwargs={'pk': main_submission.id})
            else:
                success_url = reverse('contests:sub-submission-create',
                                      kwargs={'problem_id': main_submission.problem.id,
                                              'submission_id': main_submission.id})
        else:
            success_url = super().get_success_url()
        if 'from_assignment' in self.request.GET:
            success_url += '?from_assignment=1'
        return success_url


class SubmissionUpdate(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, UpdateView):
    model = Submission
    form_class = SubmissionUpdateForm
    template_name = 'contests/submission/submission_update_form.html'
    permission_required = 'contests.change_submission'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['from_assignment'] = 'from_assignment' in self.request.GET
        return context

    def get_success_url(self):
        if self.object.main_submission is not None and 'from_main_submission' in self.request.GET:
            success_url = self.object.main_submission.get_absolute_url()
        else:
            success_url = self.object.get_absolute_url()
        if 'from_assignment' in self.request.GET:
            success_url += '?from_assignment=1'
        return success_url


class SubmissionUpdateAPI(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                          BaseUpdateView):
    model = Submission
    form_class = SubmissionUpdateForm
    http_method_names = ['post']
    permission_required = 'contests.change_submission'

    def http_method_not_allowed(self, request, *args, **kwargs):
        return JsonResponse({'status': 'http_method_not_allowed'})

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def form_valid(self, form):
        self.object = form.save()
        response = {
            'status': 'ok',
            'updated': [
                {
                    'id': self.object.id,
                    'status': self.object.status,
                    'status_display': self.object.get_status_display(),
                    'score': self.object.score,
                    'color': colorize(self.object.status)
                }
            ]
        }
        if self.object.main_submission is not None:
            response['updated'].append({
                'id': self.object.main_submission.id,
                'status': self.object.main_submission.status,
                'status_display': self.object.main_submission.get_status_display(),
                'score': self.object.main_submission.score,
                'color': colorize(self.object.main_submission.status)
            })
        Notification.objects.notify(self.object.owner, subject=self.request.user, action="изменил оценку Вашей посылки",
                                    object=self.object, relation="из раздела", reference=self.object.contest)
        return JsonResponse(response)

    def form_invalid(self, form):
        return JsonResponse({'status': 'form_errors', 'errors': form.errors})

    def handle_no_permission(self):
        return JsonResponse({'status': 'access_denied'})


class SubmissionMoss(LoginRedirectMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin,
                     SingleObjectMixin, FormView):
    model = Submission
    form_class = SubmissionMossForm
    template_name = 'contests/submission/submission_moss_form.html'
    permission_required = 'contests.moss_submission'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        return super().get(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        return super().post(request, *args, **kwargs)

    def form_valid(self, form):
        to_submission_ids = list(form.cleaned_data['to_submissions'].values_list('id', flat=True))
        self.object.moss_to_submissions = ','.join(map(str, to_submission_ids))
        self.object.moss_report_url = None
        self.object.save()
        moss_submission.delay(self.object.id, to_submission_ids)
        return HttpResponseRedirect(self.get_success_url())

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['submission'] = self.object
        return kwargs

    def get_success_url(self):
        return reverse('contests:submission-detail', kwargs={'pk': self.kwargs['pk']})


class SubmissionEvaluateAPI(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, View):
    http_method_names = ['get']
    permission_required = 'contests.evaluate_submission'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['sandbox_type'] = request.GET.get('sandbox', 'subprocess')
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def http_method_not_allowed(self, request, *args, **kwargs):
        return JsonResponse({'status': 'http_method_not_allowed'})

    def get_object(self):
        return Submission.objects.get(pk=self.kwargs.get('pk'))

    def get(self, request, *args, **kwargs):
        submission = self.get_object()
        if submission.problem.type == 'Program' and submission.problem.is_testable:
            task = evaluate_submission.delay(submission.pk, request.user.id, **self.storage)
            submission.task_id = task.id
            submission.save()
            return JsonResponse({'status': 'task_queued', 'task_id': task.id})
        return JsonResponse({'status': 'ignored'})

    def handle_no_permission(self):
        return JsonResponse({'status': 'access_denied'})


class SubmissionProgressAPI(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, View):
    http_method_names = ['get']
    permission_required = 'contests.evaluate_submission'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def http_method_not_allowed(self, request, *args, **kwargs):
        return JsonResponse({'status': 'http_method_not_allowed'})

    def get_object(self):
        return Submission.objects.get(pk=self.kwargs.get('pk'))

    def get(self, request, *args, **kwargs):
        task_id = self.kwargs.get('task_id')
        progress = TaskProgress(task_id)
        return JsonResponse(progress.get_info())

    def handle_no_permission(self):
        return JsonResponse({'status': 'access_denied'})


class SubmissionClearTaskAPI(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, View):
    http_method_names = ['get']
    permission_required = 'contests.evaluate_submission'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def http_method_not_allowed(self, request, *args, **kwargs):
        return JsonResponse({'status': 'http_method_not_allowed'})

    def get_object(self):
        return Submission.objects.get(pk=self.kwargs.get('pk'))

    def get(self, request, *args, **kwargs):
        submission = self.get_object()
        submission.task_id = None
        submission.save()
        return JsonResponse({'status': 'task_cleared'})

    def handle_no_permission(self):
        return JsonResponse({'status': 'access_denied'})


class SubmissionDelete(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, DeleteView):
    model = Submission
    template_name = 'contests/submission/submission_delete.html'
    permission_required = 'contests.delete_submission'

    def has_ownership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.owner_id == self.request.user.id

    def has_leadership(self):
        if not hasattr(self, 'object'):
            self.object = self.get_object()
        return self.object.course.leaders.filter(id=self.request.user.id).exists()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['from_assignment'] = 'from_assignment' in self.request.GET
        return context

    def get_success_url(self):
        if self.object.main_submission is not None and 'from_main_submission' in self.request.GET:
            success_url = self.object.main_submission.get_absolute_url()
        elif self.object.assignment is not None and 'from_assignment' in self.request.GET:
            success_url = self.object.assignment.get_absolute_url()
        else:
            success_url = self.object.problem.get_absolute_url()
        if 'from_assignment' in self.request.GET:
            success_url += '?from_assignment=1'
        return success_url


class SubmissionList(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, ListView):
    model = Submission
    template_name = 'contests/submission/submission_list.html'
    context_object_name = 'submissions'
    permission_required = 'contests.view_submission_list'
    paginate_by = 20

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        course_id = self.kwargs.get('course_id', None)
        if course_id:
            self.storage['course'] = get_object_or_404(Course, id=course_id)
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return 'course' in self.storage and self.storage['course'].owner_id == self.request.user.id

    def has_leadership(self):
        return 'course' in self.storage and self.storage['course'].leaders.filter(id=self.request.user.id).exists()

    def get_queryset(self):
        queryset = super().get_queryset().select_related('owner', 'problem', 'problem__contest')
        if 'course' in self.storage:
            queryset = queryset.filter(problem__contest__course_id=self.storage['course'].id)
        else:
            course_ids = self.request.user.filter_set.values_list('course_id')
            queryset = (queryset.filter(problem__contest__course_id__in=course_ids)
                        .filter(problem__contest__course__faculty=self.request.user.account.faculty)
                        .select_related('problem__contest__course'))
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(self.storage)
        return context


class SubmissionBackup(LoginRedirectMixin, PermissionRequiredMixin, BaseListView):
    model = Submission
    permission_required = 'contests.download_submission'

    def dispatch(self, request, *args, **kwargs):
        course_id = self.kwargs.get('course_id', None)
        if not Course.objects.filter(id=course_id).exists():
            raise Http404("Course does not exist")
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, *args, **kwargs):
        self.object_list = self.get_queryset()
        course_id = self.kwargs.get('course_id', None)
        response = HttpResponse(Submission.objects.backup(self.object_list), content_type='application/zip')
        response['Content-Disposition'] = 'attachment; filename={}.zip'.format(course_id)
        return response

    def get_queryset(self):
        course_id = self.kwargs.get('course_id', None)
        contest_id = int(self.request.GET.get('contest_id') or 0)
        year = int(self.request.GET.get('year') or 0)
        queryset = super().get_queryset().filter(problem__contest__course_id=course_id).select_related('problem')
        if contest_id:
            queryset = queryset.filter(problem__contest__id=contest_id)
        if year:
            queryset = queryset.filter(date_created__year__gte=year)
        return queryset


"""=================================================== Execution ===================================================="""


class ExecutionList(LoginRequiredMixin, LeadershipOrMixin, OwnershipOrMixin, PermissionRequiredMixin, ListView):
    model = Execution
    template_name = 'contests/execution/execution_list.html'
    context_object_name = 'executions'
    permission_required = 'contests.evaluate_submission'

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.storage = dict()

    def dispatch(self, request, *args, **kwargs):
        self.storage['submission'] = get_object_or_404(Submission, id=kwargs.get('pk'))
        return super().dispatch(request, *args, **kwargs)

    def has_ownership(self):
        return self.storage['submission'].owner_id == self.request.user.id

    def has_leadership(self):
        return self.storage['submission'].course.leaders.filter(id=self.request.user.id).exists()

    def get_queryset(self):
        return Execution.objects.filter(submission_id=self.storage['submission'].id)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['course'] = self.storage['submission'].course
        return context


"""==================================================== Specific ===================================================="""


@login_required
def index(request):
    notifications = Notification.objects.get_queryset().filter(recipient=request.user).unread()[:10]
    schedules = Schedule.objects.get_queryset().filter(
        date_from__range=[timezone.now().date() - timedelta(days=6), timezone.now().date() + timedelta(days=7)])
    announcements = Announcement.objects.get_queryset().filter(
        date_created__range=[timezone.now().date() - timedelta(days=90), timezone.now().date()])
    count_of_news = notifications.count() + schedules.count() + announcements.count()
    has_unread_notifications = Notification.objects.get_queryset().filter(
        recipient=request.user).unread().exists()
    if request.user.has_perm('contests.add_course'):
        filtered_course_ids = request.user.filter_set.values_list('course_id')
        filtered_courses = Course.objects.filter(id__in=filtered_course_ids)
        return render(request, 'contests/index.html',
                      {'courses': filtered_courses, 'notifications': notifications, 'schedules': schedules,
                       'announcements': announcements, 'count_of_news': count_of_news,
                       'has_unread_notifications': has_unread_notifications})
    else:
        return redirect(reverse('contests:assignment-list'))
